from typing import Any, BinaryIO, Dict, List
from langchain_core.documents import Document
from pypdf import PdfReader
from .base import BaseParser
import docx

class PDFParser(BaseParser):
    def parse(self, file_obj: BinaryIO, filename: str, **kwargs: Any) -> Dict[str, Any]:
        reader = PdfReader(file_obj)
        documents: List[Document] = []

        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            if not text.strip():
                continue

            doc = Document(
                page_content=text.strip(),
                metadata={
                    "page_number": page_num,
                    "total_pages": len(reader.pages),
                },
            )
            documents.append(doc)

        return {
            "content": documents,
            "metadata": {
                "filename": filename,
                "type": "texte",
                "total_chunks": len(documents),
            },
        }


class WordParser(BaseParser):
    def parse(self, file_obj: BinaryIO, filename: str, **kwargs: Any) -> Dict[str, Any]:
        doc = docx.Document(file_obj)
        documents: List[Document] = []

        for idx, para in enumerate(doc.paragraphs, start=1):
            text = para.text.strip()

            if not text:
                continue

            document = Document(
                page_content=text,
                metadata={
                    "paragraph_number": idx,
                    "total_paragraphs": len(doc.paragraphs),
                },
            )
            documents.append(document)

        return {
            "content": documents,
            "metadata": {
                "filename": filename,
                "type": "texte",
                "total_chunks": len(documents),
            },
        }